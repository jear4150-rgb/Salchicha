import csv
import io
import unicodedata

from openpyxl import load_workbook

ADDRESS_KEYWORDS = ["direccion", "domicilio", "address", "calle", "ubicacion"]
CLIENTE_KEYWORDS = ["cliente", "nombre", "name", "customer"]
TELEFONO_KEYWORDS = ["telefono", "phone", "celular", "cel"]


def _strip_accents(text: str) -> str:
    normalized = unicodedata.normalize("NFKD", text)
    return "".join(c for c in normalized if not unicodedata.combining(c))


def _find_column(headers: list[str], keywords: list[str], exclude: set[str]) -> str | None:
    for header in headers:
        if header in exclude:
            continue
        normalized = _strip_accents(header.lower())
        if any(kw in normalized for kw in keywords):
            return header
    return None


def _decode_csv_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_csv(data: bytes) -> list[dict[str, str]]:
    text = _decode_csv_bytes(data)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = []
    for row in reader:
        cleaned = {(k or "").strip(): (v or "").strip() for k, v in row.items() if k}
        if any(cleaned.values()):
            rows.append(cleaned)
    return rows


def _parse_xlsx(data: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    all_rows = list(sheet.iter_rows(values_only=True))
    if not all_rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in all_rows[0]]
    rows = []
    for raw_row in all_rows[1:]:
        row = {
            headers[i]: ("" if raw_row[i] is None else str(raw_row[i]).strip())
            for i in range(len(headers))
            if headers[i]
        }
        if any(row.values()):
            rows.append(row)
    return rows


def parse_spreadsheet(filename: str, data: bytes) -> list[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(data)
    return _parse_csv(data)


def detect_columns(headers: list[str]) -> dict[str, str | None]:
    cliente_col = _find_column(headers, CLIENTE_KEYWORDS, set())
    telefono_col = _find_column(headers, TELEFONO_KEYWORDS, set())
    exclude = {c for c in (cliente_col, telefono_col) if c}
    direccion_col = _find_column(headers, ADDRESS_KEYWORDS, exclude)

    if not direccion_col:
        candidates = [h for h in headers if h not in exclude]
        if len(candidates) == 1:
            direccion_col = candidates[0]

    return {"direccion": direccion_col, "cliente": cliente_col, "telefono": telefono_col}


def guess_address_column(headers: list[str], rows: list[dict[str, str]], exclude: set[str]) -> str | None:
    """Fallback: pick the text column with the longest average length among the candidates."""
    candidates = [h for h in headers if h not in exclude]
    if not candidates:
        return None
    best_col, best_avg = None, 0.0
    for col in candidates:
        lengths = [len(row.get(col, "")) for row in rows if row.get(col)]
        if not lengths:
            continue
        avg = sum(lengths) / len(lengths)
        if avg > best_avg:
            best_col, best_avg = col, avg
    return best_col
